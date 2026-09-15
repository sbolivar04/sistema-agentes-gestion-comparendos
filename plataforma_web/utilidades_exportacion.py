"""
Módulo de utilidades para generación de reportes y exportación en Excel (.xlsx).
Utiliza openpyxl con estilos corporativos, formatos de moneda y ajuste dinámico de columnas.
"""

import io
from datetime import datetime
from typing import List, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# Estilos Corporativos
COLOR_ENCABEZADO = "0F172A"       # Azul oscuro / Slate corporativo
COLOR_TEXTO_ENCABEZADO = "FFFFFF"
COLOR_FILA_PAR = "F8FAFC"
COLOR_TOTALES = "F1F5F9"

FUENTE_TITULO = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
FUENTE_SUBTITULO = Font(name="Segoe UI", size=9, italic=True, color="64748B")
FUENTE_ENCABEZADO = Font(name="Segoe UI", size=10, bold=True, color=COLOR_TEXTO_ENCABEZADO)
FUENTE_DATOS = Font(name="Segoe UI", size=9)
FUENTE_TOTALES = Font(name="Segoe UI", size=9, bold=True)

FILL_ENCABEZADO = PatternFill(start_color=COLOR_ENCABEZADO, end_color=COLOR_ENCABEZADO, fill_type="solid")
FILL_PAR = PatternFill(start_color=COLOR_FILA_PAR, end_color=COLOR_FILA_PAR, fill_type="solid")
FILL_TOTALES = PatternFill(start_color=COLOR_TOTALES, end_color=COLOR_TOTALES, fill_type="solid")

BORDE_FINO = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0")
)
BORDE_TOTALES = Border(
    top=Side(style="thin", color="94A3B8"),
    bottom=Side(style="double", color="0F172A")
)

FORMATO_MONEDA = '"$"#,##0'
FORMATO_FECHA = "yyyy-mm-dd"


def _formatear_fecha(dt) -> str:
    if not dt:
        return "N/A"
    if isinstance(dt, str):
        return dt.split("T")[0].split(" ")[0]
    return dt.strftime("%Y-%m-%d")


def generar_excel_resumen(comparendos_serializados: List[dict]) -> io.BytesIO:
    """
    Genera un archivo Excel (.xlsx) con el reporte ejecutivo y operativo resumido.
    Contiene las columnas esenciales para la gestión de flota y liquidación de pagos.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen Comparendos"
    ws.views.sheetView[0].showGridLines = True

    # 1. Título y Metadatos del Reporte
    ws.merge_cells("A1:Q1")
    ws["A1"] = "GESTIÓN INTELIGENTE DE COMPARENDOS SIMIT - REPORTE RESUMEN"
    ws["A1"].font = FUENTE_TITULO
    ws["A1"].alignment = Alignment(vertical="center")

    ws.merge_cells("A2:Q2")
    ws["A2"] = f"Generado el: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} • Total Registros: {len(comparendos_serializados)}"
    ws["A2"].font = FUENTE_SUBTITULO
    ws["A2"].alignment = Alignment(vertical="center")

    # 2. Encabezados de Columnas
    encabezados = [
        "Placa",
        "No. Comparendo",
        "No. Resolución",
        "Tipo Registro",
        "Código Infracción",
        "Descripción Infracción",
        "Secretaría de Tránsito",
        "Fecha Infracción",
        "Fecha Notificación",
        "Valor Nominal ($)",
        "Intereses Mora ($)",
        "Valor Total ($)",
        "Beneficio Ley",
        "Fecha Límite Descuento",
        "Total a Pagar Hoy ($)",
        "Ahorro Disponible ($)",
        "Estado SIMIT"
    ]

    fila_encabezado = 4
    for col_idx, texto in enumerate(encabezados, start=1):
        celda = ws.cell(row=fila_encabezado, column=col_idx, value=texto)
        celda.font = FUENTE_ENCABEZADO
        celda.fill = FILL_ENCABEZADO
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = BORDE_FINO
    ws.row_dimensions[fila_encabezado].height = 28

    # 3. Filas de Datos
    fila_actual = 5
    for idx, c in enumerate(comparendos_serializados):
        val_nominal = float(c.get("valor_nominal") or 0)
        intereses = float(c.get("intereses") or 0)
        val_total = float(c.get("valor_total") or 0)
        val_pagar = float(c.get("valor_a_pagar") or val_total)
        ahorro = float(c.get("ahorro_disponible") or 0)

        fila_datos = [
            c.get("placa") or "N/A",
            c.get("numero_comparendo") or "N/A",
            c.get("numero_resolucion") or "N/A",
            c.get("tipo_registro") or "Comparendo",
            c.get("codigo_infraccion") or "",
            c.get("descripcion_infraccion") or "",
            c.get("secretaria") or "",
            _formatear_fecha(c.get("fecha_infraccion")),
            _formatear_fecha(c.get("fecha_notificacion")),
            val_nominal,
            intereses,
            val_total,
            c.get("etiqueta_descuento") or "Sin Descuento",
            c.get("fecha_limite_descuento") or "Vencido",
            val_pagar,
            ahorro,
            c.get("estado_simit") or "Activo"
        ]

        usar_par = (idx % 2 == 1)
        for col_idx, valor in enumerate(fila_datos, start=1):
            celda = ws.cell(row=fila_actual, column=col_idx, value=valor)
            celda.font = FUENTE_DATOS
            celda.border = BORDE_FINO
            if usar_par:
                celda.fill = FILL_PAR

            # Alineaciones y formatos numéricos
            if col_idx in (10, 11, 12, 15, 16):  # Columnas de moneda
                celda.number_format = FORMATO_MONEDA
                celda.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in (1, 4, 5, 8, 9, 14, 17):  # Centradas
                celda.alignment = Alignment(horizontal="center", vertical="center")
            else:
                celda.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[fila_actual].height = 20
        fila_actual += 1

    # 4. Fila de Totales
    if len(comparendos_serializados) > 0:
        ws.cell(row=fila_actual, column=1, value="TOTALES GENERALES").font = FUENTE_TOTALES
        ws.cell(row=fila_actual, column=1).alignment = Alignment(horizontal="left", vertical="center")

        for col_idx in range(1, len(encabezados) + 1):
            celda = ws.cell(row=fila_actual, column=col_idx)
            celda.fill = FILL_TOTALES
            celda.border = BORDE_TOTALES
            celda.font = FUENTE_TOTALES

            col_letra = get_column_letter(col_idx)
            if col_idx in (10, 11, 12, 15, 16):
                celda.value = f"=SUM({col_letra}5:{col_letra}{fila_actual - 1})"
                celda.number_format = FORMATO_MONEDA
                celda.alignment = Alignment(horizontal="right", vertical="center")

        ws.row_dimensions[fila_actual].height = 24

    # 5. Autoajuste dinámico del ancho de columnas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 4:  # Saltar títulos combinados
                continue
            val_str = str(cell.value or "")
            if val_str.startswith("="):
                val_str = "$999,999,999"
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generar_excel_detallado(comparendos_serializados: List[dict]) -> io.BytesIO:
    """
    Genera un archivo Excel (.xlsx) exhaustivo con toda la información técnica,
    jurídica y operativa de los comparendos y multas.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Detalle Comparendos"
    ws.views.sheetView[0].showGridLines = True

    # 1. Título y Metadatos
    ws.merge_cells("A1:X1")
    ws["A1"] = "GESTIÓN INTELIGENTE DE COMPARENDOS SIMIT - REPORTE COMPLETO Y DETALLADO"
    ws["A1"].font = FUENTE_TITULO
    ws["A1"].alignment = Alignment(vertical="center")

    ws.merge_cells("A2:X2")
    ws["A2"] = f"Generado el: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} • Auditoría Completa de Flota"
    ws["A2"].font = FUENTE_SUBTITULO
    ws["A2"].alignment = Alignment(vertical="center")

    # 2. Encabezados de Columnas
    encabezados = [
        "ID Sistema",
        "Criterio / NIT Base",
        "Placa Vehicular",
        "No. Comparendo",
        "No. Resolución",
        "Tipo Registro",
        "Código Infracción",
        "Descripción Oficial Infracción",
        "Secretaría de Tránsito",
        "Dirección del Hecho",
        "Fuente Comparendo",
        "Es Fotodetección",
        "Fecha Infracción",
        "Fecha Notificación",
        "Fecha Resolución",
        "Valor Nominal SIMIT ($)",
        "Intereses de Mora ($)",
        "Valor Total SIMIT ($)",
        "Beneficio / Descuento",
        "Fecha Límite Descuento",
        "Total a Pagar Hoy ($)",
        "Ahorro Legal Potencial ($)",
        "Estado SIMIT"
    ]

    fila_encabezado = 4
    for col_idx, texto in enumerate(encabezados, start=1):
        celda = ws.cell(row=fila_encabezado, column=col_idx, value=texto)
        celda.font = FUENTE_ENCABEZADO
        celda.fill = FILL_ENCABEZADO
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = BORDE_FINO
    ws.row_dimensions[fila_encabezado].height = 28

    # 3. Filas de Datos
    fila_actual = 5
    for idx, c in enumerate(comparendos_serializados):
        val_nominal = float(c.get("valor_nominal") or 0)
        intereses = float(c.get("intereses") or 0)
        val_total = float(c.get("valor_total") or 0)
        val_pagar = float(c.get("valor_a_pagar") or val_total)
        ahorro = float(c.get("ahorro_disponible") or 0)
        es_foto = "Sí (Cámara)" if c.get("es_fotodeteccion") else "No (Físico en vía)"

        fila_datos = [
            c.get("id") or idx + 1,
            c.get("criterio_busqueda") or "N/A",
            c.get("placa") or "N/A",
            c.get("numero_comparendo") or "N/A",
            c.get("numero_resolucion") or "N/A",
            c.get("tipo_registro") or "Comparendo",
            c.get("codigo_infraccion") or "",
            c.get("descripcion_infraccion") or "",
            c.get("secretaria") or "",
            c.get("direccion") or "No reportada",
            c.get("fuente_comparendo") or "No reportada",
            es_foto,
            _formatear_fecha(c.get("fecha_infraccion")),
            _formatear_fecha(c.get("fecha_notificacion")),
            _formatear_fecha(c.get("fecha_resolucion")),
            val_nominal,
            intereses,
            val_total,
            c.get("etiqueta_descuento") or "Sin Descuento",
            c.get("fecha_limite_descuento") or "Vencido",
            val_pagar,
            ahorro,
            c.get("estado_simit") or "Activo"
        ]

        usar_par = (idx % 2 == 1)
        for col_idx, valor in enumerate(fila_datos, start=1):
            celda = ws.cell(row=fila_actual, column=col_idx, value=valor)
            celda.font = FUENTE_DATOS
            celda.border = BORDE_FINO
            if usar_par:
                celda.fill = FILL_PAR

            # Formatos de moneda en columnas 16, 17, 18, 21, 22
            if col_idx in (16, 17, 18, 21, 22):
                celda.number_format = FORMATO_MONEDA
                celda.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in (1, 2, 3, 6, 7, 12, 13, 14, 15, 19, 20, 23):
                celda.alignment = Alignment(horizontal="center", vertical="center")
            else:
                celda.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[fila_actual].height = 20
        fila_actual += 1

    # 4. Fila de Totales
    if len(comparendos_serializados) > 0:
        ws.cell(row=fila_actual, column=1, value="TOTALES GENERALES").font = FUENTE_TOTALES
        ws.cell(row=fila_actual, column=1).alignment = Alignment(horizontal="left", vertical="center")

        for col_idx in range(1, len(encabezados) + 1):
            celda = ws.cell(row=fila_actual, column=col_idx)
            celda.fill = FILL_TOTALES
            celda.border = BORDE_TOTALES
            celda.font = FUENTE_TOTALES

            col_letra = get_column_letter(col_idx)
            if col_idx in (16, 17, 18, 21, 22):
                celda.value = f"=SUM({col_letra}5:{col_letra}{fila_actual - 1})"
                celda.number_format = FORMATO_MONEDA
                celda.alignment = Alignment(horizontal="right", vertical="center")

        ws.row_dimensions[fila_actual].height = 24

    # 5. Autoajuste de Ancho
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 4:
                continue
            val_str = str(cell.value or "")
            if val_str.startswith("="):
                val_str = "$999,999,999"
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
